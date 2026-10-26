import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from PIL import Image, ImageOps
import openpyxl
from openpyxl.chart import ScatterChart, Reference, Series

# ==============================================================================
# CONFIGURACIÓN DE ARCHIVOS Y CASOS
# ==============================================================================
CASOS = [
    {
        "id": 1,
        "nombre": "Empotrado - Empotrado",
        "hoja": "1_Empotrado_Empotrado",
        "archivo": "1. Empotrado - empotrado.jpg",
        "instruccion": "Deja la PRENSA BLANCA a la izquierda (x=0), LIBROS a la derecha (x=L) y pachón ABAJO.",
        # Fórmula Excel teórica para x en B{row} (con carga en x = 2/3)
        "formula_excel": "=6*B{r}^2 - 7*B{r}^3 + IF(B{r}>=2/3, 27*(B{r}-2/3)^3, 0)",
        "func_teorica": lambda x: 6*x**2 - 7*x**3 + np.where(x >= 2/3, 27*(x - 2/3)**3, 0.0)
    },
    {
        "id": 2,
        "nombre": "Abisagrado - Abisagrado",
        "hoja": "2_Abisagrado_Abisagrado",
        "archivo": "2. Abisagrado - abisagrado.jpg",
        "instruccion": "Deja el LÁPIZ IZQUIERDO (silla) en x=0, LÁPIZ DERECHO (libros) en x=L y pachón ABAJO.",
        "formula_excel": "=8*B{r} - 9*B{r}^3 + IF(B{r}>=2/3, 27*(B{r}-2/3)^3, 0)",
        "func_teorica": lambda x: 8*x - 9*x**3 + np.where(x >= 2/3, 27*(x - 2/3)**3, 0.0)
    },
    {
        "id": 3,
        "nombre": "Empotrado - Abisagrado",
        "hoja": "3_Empotrado_Abisagrado",
        "archivo": "3. Empotrado - abisagrado.jpg",
        "instruccion": "Deja los LIBROS a la izquierda (x=0), LÁPIZ en la silla a la derecha (x=L) y pachón ABAJO.",
        "formula_excel": "=12*B{r}^2 - 13*B{r}^3 + IF(B{r}>=2/3, 27*(B{r}-2/3)^3, 0)",
        "func_teorica": lambda x: 12*x**2 - 13*x**3 + np.where(x >= 2/3, 27*(x - 2/3)**3, 0.0)
    },
    {
        "id": 4,
        "nombre": "Empotrado - Libre",
        "hoja": "4_Empotrado_Libre",
        "archivo": "4. Empotrado - libre.jpg",
        "instruccion": "Deja los LIBROS a la izquierda (x=0), PUNTA LIBRE a la derecha (x=L) y pachón ABAJO.\n"
                    "NOTA: Para el Clic 2 (fin del eje horizontal y=0), haz clic a la altura donde saldría la regla recta sin doblarse.",
        "formula_excel": "=2*B{r}^2 - B{r}^3 + IF(B{r}>=2/3, (B{r}-2/3)^3, 0)",
        "func_teorica": lambda x: 2*x**2 - x**3 + np.where(x >= 2/3, (x - 2/3)**3, 0.0)
    }
]


class ExtractorViga:
    def __init__(self, ruta_img, titulo, instruccion):
        # Cargar respetando rotación EXIF original
        pil_img = Image.open(ruta_img)
        pil_img = ImageOps.exif_transpose(pil_img)
        self.img = np.array(pil_img)
        self.titulo = titulo
        self.instruccion = instruccion
        self.puntos_eje = []     # 2 clics: (x=0, y=0) y (x=L, y=0)
        self.puntos_curva = []   # Clics a lo largo de la viga deformada

        self.fig, self.ax = plt.subplots(figsize=(12, 8))
        self.fig.canvas.mpl_connect('button_press_event', self.on_click)
        self.fig.canvas.mpl_connect('key_press_event', self.on_key)
        self.actualizar_vista()
        plt.show()

    def actualizar_vista(self):
        self.ax.clear()
        self.ax.imshow(self.img)
        self.ax.set_title(
            f"{self.titulo}\n{self.instruccion}\n"
            "[Tecla 'R': Rotar 90° | Tecla 'F': Espejo Horizontal | Tecla 'Z': Borrar último punto | ENTER: Terminar]",
            fontsize=10, fontweight='bold'
        )

        # Dibujar puntos de referencia del eje X (rojo)
        if len(self.puntos_eje) >= 1:
            p0 = self.puntos_eje[0]
            self.ax.plot(p0[0], p0[1], 'ro', markersize=8, label='Inicio (x=0, y=0)')
        if len(self.puntos_eje) == 2:
            p0, p1 = self.puntos_eje[0], self.puntos_eje[1]
            self.ax.plot(p1[0], p1[1], 'rs', markersize=8, label='Referencia Fin Eje (x=L, y=0)')
            self.ax.plot([p0[0], p1[0]], [p0[1], p1[1]], 'r--', linewidth=1.5, label='Eje Horizontal (y=0)')

        # Dibujar puntos marcados sobre la viga (cyan)
        if len(self.puntos_curva) > 0:
            pts = np.array(self.puntos_curva)
            self.ax.plot(pts[:, 0], pts[:, 1], 'co-', markersize=5, linewidth=1.2,
                        label=f'Puntos viga ({len(self.puntos_curva)})')

        # Mostrar qué clic toca hacer abajo
        if len(self.puntos_eje) == 0:
            msg = "PASO 1: Haz CLIC IZQUIERDO en el inicio de la viga (x = 0, extremo izquierdo)."
        elif len(self.puntos_eje) == 1:
            msg = "PASO 2: Haz CLIC IZQUIERDO en el final del eje horizontal sin deformar (x = L, y = 0)."
        else:
            msg = f"PASO 3: Haz clics a lo largo de la viga desde x=0 hasta x=L (Llevas {len(self.puntos_curva)} puntos). Presiona ENTER al terminar."

        self.ax.set_xlabel(msg, fontsize=11, color='blue', fontweight='bold')
        self.ax.legend(loc='upper right')
        self.fig.canvas.draw()

    def on_click(self, event):
        # Ignorar clics si se está usando el zoom o la lupa de la barra de matplotlib
        if event.inaxes != self.ax or self.fig.canvas.toolbar.mode != '':
            return
        if event.button == 1:  # Clic izquierdo
            if len(self.puntos_eje) < 2:
                self.puntos_eje.append((event.xdata, event.ydata))
                # Si acaba de marcar x=0, también lo agregamos como primer punto de la viga
                if len(self.puntos_eje) == 1:
                    self.puntos_curva.append((event.xdata, event.ydata))
            else:
                self.puntos_curva.append((event.xdata, event.ydata))
            self.actualizar_vista()

    def on_key(self, event):
        if event.key in ['r', 'R']:
            self.img = np.rot90(self.img, -1)
            self.puntos_eje.clear()
            self.puntos_curva.clear()
            self.actualizar_vista()
        elif event.key in ['f', 'F']:
            self.img = np.fliplr(self.img)
            self.puntos_eje.clear()
            self.puntos_curva.clear()
            self.actualizar_vista()
        elif event.key in ['z', 'Z', 'backspace']:
            if len(self.puntos_curva) > 1:
                self.puntos_curva.pop()
            elif len(self.puntos_eje) > 0:
                self.puntos_eje.pop()
                self.puntos_curva.clear()
            self.actualizar_vista()
        elif event.key == 'enter':
            if len(self.puntos_eje) == 2 and len(self.puntos_curva) >= 5:
                plt.close(self.fig)
            else:
                print("Debes marcar los 2 puntos del eje y al menos 5 puntos sobre la curva antes de presionar ENTER.")

    def obtener_datos_interpolados(self):
        """
        Proyecta los píxeles marcados sobre el eje definido por (p0 -> p1),
        normaliza X en [0, 1] e interpola en los 21 puntos (0, 0.05, ..., 1.0).
        """
        p0 = np.array(self.puntos_eje[0], dtype=float)
        p1 = np.array(self.puntos_eje[1], dtype=float)
        vec_eje = p1 - p0
        L_pix = np.linalg.norm(vec_eje)
        u_x = vec_eje / L_pix
        # En imágenes, el eje Y de píxeles crece hacia abajo: vector perpendicular hacia abajo
        u_y = np.array([-u_x[1], u_x[0]])
        if u_y[1] < 0:
            u_y = -u_y

        pts = np.array(self.puntos_curva, dtype=float)
        diffs = pts - p0
        x_norm_raw = (diffs @ u_x) / L_pix
        y_pix_raw = np.abs(diffs @ u_y)

        # Ordenar por X creciente y asegurar extremos en x=0 y x=1
        orden = np.argsort(x_norm_raw)
        x_norm_raw = x_norm_raw[orden]
        y_pix_raw = y_pix_raw[orden]

        # Malla estándar de 21 puntos (0.00, 0.05, 0.10, ..., 1.00)
        x_grid = np.round(np.linspace(0.0, 1.0, 21), 2)
        y_exp_interp = np.interp(x_grid, x_norm_raw, y_pix_raw)
        y_exp_interp[0] = 0.0  # Condición en x=0
        return x_grid, y_exp_interp


def ejecutar_proyecto():
    wb = openpyxl.Workbook()
    # Eliminar hoja por defecto al final
    hoja_default = wb.active
    resumen_rmse = []

    for caso in CASOS:
        ruta = caso["archivo"]
        if not os.path.exists(ruta):
            print(f"[AVISO] No se encontró '{ruta}' en esta carpeta. Saltando caso...")
            continue

        print(f"\n---> Procesando Caso {caso['id']}: {caso['nombre']}...")
        extractor = ExtractorViga(ruta, f"Caso {caso['id']}: {caso['nombre']}", caso["instruccion"])
        x_vals, y_exp = extractor.obtener_datos_interpolados()

        # En los casos 1, 2 y 3 el extremo derecho x=1 también tiene y(1)=0
        if caso["id"] in [1, 2, 3]:
            y_exp[-1] = 0.0

        y_teo = caso["func_teorica"](x_vals)
        y_teo_norm = y_teo / np.max(y_teo)
        y_exp_norm = y_exp / np.max(y_exp) if np.max(y_exp) > 0 else y_exp

        # Cálculo de RMSE sobre las curvas normalizadas
        rmse = np.sqrt(np.mean((y_exp_norm - y_teo_norm) ** 2))
        resumen_rmse.append((caso["nombre"], rmse))
        print(f"     RMSE ({caso['nombre']}): {rmse:.6f}")

        # 1. Guardar gráfico individual en PNG de alta calidad
        plt.figure(figsize=(8, 5), dpi=150)
        plt.plot(x_vals, y_teo_norm, 'o-', color='#2b5c8f', label='Y Teórica Normalizada', linewidth=2)
        plt.plot(x_vals, y_exp_norm, 'o-', color='#d95f02', label='Y Experimental Normalizada', linewidth=2)
        plt.title(f"CURVA DE FLEXIÓN: {caso['nombre'].upper()} (CARGA EN x=2L/3)", fontsize=11, fontweight='bold')
        plt.xlabel("Posición normalizada X (x/L)")
        plt.ylabel("Deflexión normalizada (y / y_max)")
        plt.xlim(0, 1.05)
        plt.ylim(0, 1.15)
        plt.grid(True, linestyle='--', alpha=0.6)
        plt.legend()
        plt.tight_layout()
        plt.savefig(f"Grafico_{caso['hoja']}.png")
        plt.close()

        # 2. Crear hoja en el archivo Excel con el formato exacto de la Figura #3
        ws = wb.create_sheet(title=caso["hoja"])
        encabezados = ["X", "Y Teórica", "Y Experimental", "Y Teórica Normalizada", "Y Experimental Normalizada"]
        for col_idx, texto in enumerate(encabezados, start=2):  # Columnas B a F, Fila 2
            ws.cell(row=2, column=col_idx, value=texto)

        for i in range(len(x_vals)):
            r = i + 3  # Filas 3 a 23
            ws.cell(row=r, column=2, value=float(x_vals[i]))                        # Col B: X
            ws.cell(row=r, column=3, value=caso["formula_excel"].format(r=r))       # Col C: Fórmula Y Teórica
            ws.cell(row=r, column=4, value=float(np.round(y_exp[i], 6)))            # Col D: Y Experimental
            ws.cell(row=r, column=5, value=f"=C{r}/MAX($C$3:$C$23)")                # Col E: Y Teórica Norm
            ws.cell(row=r, column=6, value=f"=D{r}/MAX($D$3:$D$23)")                # Col F: Y Exp Norm

        # Celda con el RMSE calculado con fórmula de Excel y valor numérico al lado
        ws.cell(row=25, column=2, value="RMSE (Fórmula):")
        ws.cell(row=25, column=3, value="=SQRT(SUMXMY2(F3:F23, E3:E23)/COUNT(E3:E23))")
        ws.cell(row=26, column=2, value="RMSE (Valor Python):")
        ws.cell(row=26, column=3, value=float(np.round(rmse, 6)))

        # Insertar gráfico nativo de dispersión en Excel
        chart = ScatterChart()
        chart.title = f"Curva de Flexión: {caso['nombre']} (x = 2L/3)"
        chart.x_axis.title = "X Normalizada"
        chart.y_axis.title = "Y Normalizada"
        x_ref = Reference(ws, min_col=2, min_row=3, max_row=23)
        y_teo_ref = Reference(ws, min_col=5, min_row=2, max_row=23)
        y_exp_ref = Reference(ws, min_col=6, min_row=2, max_row=23)
        s1 = Series(y_teo_ref, x_ref, title_from_data=True)
        s2 = Series(y_exp_ref, x_ref, title_from_data=True)
        chart.series.append(s1)
        chart.series.append(s2)
        ws.add_chart(chart, "H2")

    if len(wb.sheetnames) > 1:
        wb.remove(hoja_default)
    nombre_excel = "MicroProyecto3_Flexion_Vigas.xlsx"
    wb.save(nombre_excel)

    print("\n" + "="*55)
    print(f"¡PROCESO COMPLETADO! Archivo guardado: {nombre_excel}")
    print("="*55)
    for nombre, val in resumen_rmse:
        print(f" - {nombre:25s} -> RMSE = {val:.5f}")


if __name__ == "__main__":
    ejecutar_proyecto()